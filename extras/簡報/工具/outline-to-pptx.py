# -*- coding: utf-8 -*-
"""大綱.md → 可編輯的 .pptx —— 標題／主張／條列都是真的文字框（而且是真的
placeholder），可以直接在 PowerPoint 或 Google Slides 裡改字、刪行、刪頁，
大綱檢視裡也看得到。

用法：
    python outline-to-pptx.py 大綱.md -o 簡報.pptx
    python outline-to-pptx.py 大綱.md -o 簡報.pptx --template 既有範本.pptx
    python outline-to-pptx.py 大綱.md -o 簡報.pptx --strict     # 檢查沒過就不出檔

只需要 python-pptx（純 Python，不需要 LibreOffice、不需要管理員權限）。
裝法見 ../../../setup/Python環境.md。流程與判準見 ../規劃.md。

⚠️ 這是**單向輸出**：在 pptx 裡改的東西不會回到大綱 .md。
   出片之後就以 pptx 為正本，大綱表退位成「當初怎麼決定的」紀錄，不要維護兩份。

⚠️ 為什麼用 placeholder 而不是自己畫文字框：placeholder 才會吃範本的母片格式、
   才會出現在大綱檢視裡、對方才改得順。自己畫的文字框在版面上看起來一樣，
   但接手的人會發現「大綱檢視是空的」——而且不會有任何錯誤訊息。
"""

import argparse
import os
import re
import sys

# Windows 主控台預設是 cp950，中文與 emoji 會變亂碼（見 Windows_powershell.md）
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, 'reconfigure'):
        _stream.reconfigure(encoding='utf-8', errors='replace')

try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit('缺少 python-pptx。安裝方式見 setup/Python環境.md：\n'
             '    python -m venv .venv\n'
             '    .venv/Scripts/pip install python-pptx\n'
             '（PyPI 有提供 SHA256，符合本機的套件驗證要求）')

# ── 大綱格式 ────────────────────────────────────────────────────────────────
# 頁標題：  ## 3 · 兩條線的節奏刻意不同     （也吃 "## 3. 標題"）
# 欄位：    | **主張** | ... |
RE_PAGE = re.compile(r'^##\s+(\d+)\s*[·.、]\s*(.+?)\s*$')
RE_ROW = re.compile(r'^\|\s*\*{0,2}([^|*]+?)\*{0,2}\s*\|\s*(.*?)\s*\|\s*$')
RE_BULLET_MARK = re.compile(r'^\s*(?:[①②③④⑤⑥⑦⑧⑨⑩]|[-*]|\d+[.)])\s*')

# 印進簡報的欄位；其餘（服務哪一題／秒／出處）只用於檢查
FIELD_TITLE, FIELD_SUB = '標題', '副標'
FIELD_CLAIM, FIELD_BODY = '主張', '內文'
FIELD_VISUAL, FIELD_NOTES = '視覺', '講稿'
FIELD_SERVES, FIELD_SEC = '服務哪一題', '秒'

MAX_BULLETS = 4          # 硬規矩一
MAX_BULLET_UNITS = 40    # 硬規矩一的「20 字」＝ 40 單位，見 width()


def width(text):
    """一點的視覺寬度，單位制：中文一字兩單位、ASCII 一單位。

    ⚠️ 不能用 len() 硬算字元數。`github.com/owner/repo` 這種純 ASCII 的字串
       字元數很多，但排起來只有一半寬 —— 用字元數會把放得下的東西誤判成超標。
       上限 40 單位就是「20 個中文字」那條規矩的正確寫法。
    """
    return sum(2 if ord(c) > 0x2E80 else 1 for c in text)
SENTENCE_MARKS = '。，、！？；：'


def clean(value):
    """去掉 markdown 粗體；只剩 <佔位符> 與標記的欄位視為沒填。

    「① <用問句寫> ② <最多兩題>」這種半填的也算沒填 —— 不然範本原封不動
    餵進來，閘門會以為那一欄填好了。
    """
    v = (value or '').replace('**', '').strip()
    if not v:
        return ''
    marks = '①②③④⑤⑥⑦⑧⑨⑩、，,。/／-｜| 	'
    probe = re.sub('<[^<>]*>', '', v)
    if probe == v:          # 沒有任何 <佔位符> → 是人自己寫的，照收
        return v            #（「服務哪一題」正常就只填 ①②，不能當成空的）
    probe = ''.join(c for c in probe if c not in marks)
    return '' if not probe else v


def split_bullets(value):
    out = []
    for part in re.split(r'<br\s*/?>|\n', value):
        part = RE_BULLET_MARK.sub('', part).strip()
        # 整段被全形括號包起來的是「給自己看的註記」，不是條列。
        # 這一條是踩過的：`（封面不放條列）` 原本會被當成條列直接印在投影片上，
        # 而且不會有任何提示。
        if not part or (part.startswith('（') and part.endswith('）')):
            continue
        out.append(part)
    return out


def parse(path):
    """回傳 (場合卡 dict, 頁面 list)。"""
    with open(path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    card, pages, section, page = {}, [], None, None
    for line in lines:
        m = RE_PAGE.match(line)
        if m:
            page = {'no': int(m.group(1)), FIELD_TITLE: m.group(2).strip()}
            pages.append(page)
            section = 'page'
            continue
        if line.startswith('## '):
            section = 'card' if '場合卡' in line else None
            page = None
            continue
        m = RE_ROW.match(line)
        if not m:
            continue
        key, value = m.group(1).strip(), clean(m.group(2))
        if set(key) <= set('-: '):          # 表格分隔列
            continue
        if section == 'card':
            card[key] = value
        elif section == 'page' and page is not None:
            page[key] = value
    return card, pages


# ── 檢查（＝ SOP 第 4 節的密度檢查，能自動查的部分）──────────────────────────
CARD_FIELDS = ['對象', '時間', '要對方回答的', '形式', '不討論', '成功標準']


def check(card, pages):
    """回傳 (擋下來的問題, 提醒)。問題對應 SOP 的閘門一與閘門二。"""
    bad, warn = [], []

    missing = [f for f in CARD_FIELDS if not card.get(f)]
    if missing:
        bad.append('場合卡還有欄位沒填：' + '、'.join(missing) + '（SOP 第 2 節）')
    if card.get('要對方回答的'):
        n = len(re.findall(r'[①②③④⑤]', card['要對方回答的']))
        if n > 2:
            warn.append('場合卡「要對方回答的」有 %d 題 —— 超過兩題代表這其實是兩場會' % n)

    total_sec = 0
    for p in pages:
        tag = 'p%d %s' % (p['no'], p.get(FIELD_TITLE, ''))
        claim = p.get(FIELD_CLAIM, '')
        if not claim:
            bad.append('%s：沒有主張' % tag)
        elif len(claim) <= 12 and not any(c in claim for c in SENTENCE_MARKS):
            warn.append('%s：主張疑似名詞短語（「%s」）—— 改寫成完整句子（SOP 3.1）' % (tag, claim))
        if not p.get(FIELD_SERVES):
            bad.append('%s：「服務哪一題」空白 —— 砍掉或併頁（SOP 第 3 節）' % tag)

        bullets = split_bullets(p.get(FIELD_BODY, ''))
        if len(bullets) > MAX_BULLETS:
            warn.append('%s：內文 %d 點，超過 %d（硬規矩一）' % (tag, len(bullets), MAX_BULLETS))
        for b in bullets:
            if width(b) > MAX_BULLET_UNITS:
                warn.append('%s：有一點 %d 單位，超過 %d（硬規矩一，中文一字兩單位）'
                            % (tag, width(b), MAX_BULLET_UNITS))

        visual = p.get(FIELD_VISUAL, '')
        if visual and re.search(r'\.(png|jpg|jpeg|gif|svg)$', visual, re.I):
            if visual.lower().endswith('.svg'):
                warn.append('%s：pptx 不吃 SVG，先轉成 PNG' % tag)
            elif not os.path.exists(visual):
                warn.append('%s：視覺檔案找不到（%s）' % (tag, visual))

        sec = re.findall(r'\d+', p.get(FIELD_SEC, ''))
        if sec:
            total_sec += int(sec[0])

    budget = re.findall(r'\d+', card.get('時間', ''))
    if budget and total_sec:
        want = max(int(x) for x in budget) * 60
        if total_sec > want:
            warn.append('秒數加總 %d 秒，超過場合卡的 %d 秒 —— 砍頁，不要加快講速'
                        % (total_sec, want))
    if budget:
        want_pages = max(int(x) for x in budget)
        if len(pages) > want_pages + 4:
            warn.append('%d 頁 vs %s 分鐘 —— 提案場合的頁數大約等於分鐘數（SOP 第 2 節）'
                        % (len(pages), card.get('時間')))
    return bad, warn


# ── 出片 ────────────────────────────────────────────────────────────────────
SKIP_LAYOUT = ('標題投影片', 'title slide', 'section', '章節', 'blank', '空白')
PREFER_LAYOUT = ('標題及內容', '標題與內容', 'title and content')
FALLBACK_FONT = 'Microsoft JhengHei'      # 只在沒給範本時才設；有範本就吃範本的


def pick_layout(prs):
    """挑一個同時有標題與內容 placeholder 的版面配置。"""
    cands = []
    for lay in prs.slide_layouts:
        idxs = sorted(ph.placeholder_format.idx for ph in lay.placeholders)
        if 0 in idxs and len(idxs) > 1:
            cands.append((lay, [i for i in idxs if i != 0][0]))
    for want in (PREFER_LAYOUT, None):
        for lay, body in cands:
            name = lay.name.lower()
            if want and any(p in name for p in want):
                return lay, body
            if want is None and not any(s in name for s in SKIP_LAYOUT):
                return lay, body
    return cands[0] if cands else (None, None)


def set_font(run, name, size=None, bold=None):
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if not name:
        return
    run.font.name = name                       # 先設 latin，ea 才排得進正確位置
    rPr = run.font._rPr
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


# 網址：完整的 http(s)，或省略協定的常見網域（例如 github.com/owner/repo）
RE_URL = re.compile(
    r'(https?://[^\s，。、）)\]]+'
    r'|(?:[\w-]+\.)+(?:com|org|net|io|dev|sh|ai|app)/[^\s，。、）)\]]*)')


def write_text(para, text, font, size=None, bold=None):
    """把文字寫進段落；偵測到網址就切成獨立的 run 並掛上超連結。

    ⚠️ 不能用 `para.text = ...` —— 那會把整段變成**一個** run，
       掛不了個別連結。有網址時必須自己 add_run() 切段。

    省略協定的網址（`github.com/...`）補上 https:// 當連結位址，
    但**顯示文字保持原樣** —— 投影片上不該多出一截 `https://`。
    """
    para.clear()
    parts = [s for s in RE_URL.split(text) if s]
    for part in parts:
        run = para.add_run()
        run.text = part
        if RE_URL.fullmatch(part):
            run.hyperlink.address = part if part.startswith('http') else 'https://' + part
        set_font(run, font, size=size, bold=bold)


def add_visual(slide, prs, path, body, full=False):
    """圖放右欄，內文縮到左邊。回傳有沒有放成功。

    `full=True`（整頁只有圖、沒有條列）時改成整頁寬 ——
    ⚠️ 細節多的圖擠在半欄裡，節點文字會小到讀不了。
    """
    margin = Inches(0.5)
    if full:
        # ⚠️ 讓出的高度要剛好夠一行主張。多留就白費 ——
        #    這種版面的圖是被「高度」卡住的，不是寬度，多預留只會讓圖變小。
        top = body.top + Inches(0.72)
        max_w = body.width
        max_h = prs.slide_height - top - Inches(0.3)
        if max_w <= 0 or max_h <= 0:
            return False
        pic = slide.shapes.add_picture(path, body.left, top, width=max_w)
        if pic.height > max_h:
            scale = max_h / float(pic.height)
            pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
            pic.left = body.left + int((max_w - pic.width) / 2)
        return True
    left = int(prs.slide_width * 0.56)
    max_w = prs.slide_width - left - margin
    max_h = prs.slide_height - body.top - margin
    if max_w <= 0 or max_h <= 0:
        return False
    body.width = int(left - body.left - margin)
    pic = slide.shapes.add_picture(path, left, body.top, width=max_w)
    if pic.height > max_h:
        scale = max_h / float(pic.height)
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        pic.left = left + int((max_w - pic.width) / 2)
    return True


def clear_slides(prs):
    """範本裡如果已經有頁，先清掉 —— 我們要的是它的母片與版面配置，不是它的內容。"""
    id_list = prs.slides._sldIdLst
    existing = list(id_list)
    for sld_id in existing:
        prs.part.drop_rel(sld_id.get(qn('r:id')))
        id_list.remove(sld_id)
    return len(existing)


def drop_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def build(card, pages, out, template=None):
    prs = Presentation(template) if template else Presentation()
    font = None if template else FALLBACK_FONT

    dropped = clear_slides(prs)
    if dropped:
        sys.stderr.write('範本裡原有的 %d 頁已清掉 —— 只留母片與版面配置。\n' % dropped)

    ratio = 1.0
    if not template:                       # 內建範本是 4:3，撐成 16:9
        target = Inches(13.333)
        ratio = target / float(prs.slide_width)
        prs.slide_width, prs.slide_height = target, Inches(7.5)

    layout, body_idx = pick_layout(prs)
    if layout is None:
        sys.exit('範本裡找不到「有標題也有內容」的版面配置，換一個範本檔')

    for p in pages:
        slide = prs.slides.add_slide(layout)
        if ratio != 1.0:
            for ph in slide.placeholders:
                # ⚠️ 四個值要一次讀完再寫，而且四個都要寫。
                #    寫 left／width 會就地生出一個明確的 <a:xfrm>，它的 y／cy 預設是 0 ——
                #    原本從版面配置繼承來的 top／height 就這樣被清掉，標題與內文會疊在
                #    頁面最上緣。讀取順序也不能顛倒：寫完 left 之後再讀 top，拿到的已經是 0。
                #    這個壞法**不會報錯**，腳本照樣印「出了 N 頁」。
                left, top, width, height = ph.left, ph.top, ph.width, ph.height
                ph.left, ph.top = int(left * ratio), top
                ph.width, ph.height = int(width * ratio), height

        title = slide.shapes.title
        head = p.get(FIELD_TITLE, '')
        sub = p.get(FIELD_SUB, '')
        if sub:
            head = '%s — %s' % (head, sub)
        write_text(title.text_frame.paragraphs[0], head, font)

        body = slide.placeholders[body_idx]
        tf = body.text_frame
        tf.word_wrap = True
        claim, bullets = p.get(FIELD_CLAIM, ''), split_bullets(p.get(FIELD_BODY, ''))

        write_text(tf.paragraphs[0], claim, font, size=Pt(24), bold=True)
        for b in bullets:
            para = tf.add_paragraph()
            para.level = 1 if claim else 0
            write_text(para, b, font, size=Pt(18))

        visual = p.get(FIELD_VISUAL, '')
        if visual and os.path.exists(visual) and not visual.lower().endswith('.svg'):
            add_visual(slide, prs, visual, body, full=not bullets)

        notes = p.get(FIELD_NOTES, '')
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

        drop_empty_placeholders(slide)

    prs.save(out)
    return len(pages)


def main():
    ap = argparse.ArgumentParser(
        description='大綱.md → 可編輯的 .pptx。流程與判準見 extras/簡報/規劃.md')
    ap.add_argument('outline', help='大綱 .md')
    ap.add_argument('-o', '--out', default='簡報.pptx', help='輸出的 .pptx（預設 簡報.pptx）')
    ap.add_argument('--template', help='要套的範本 .pptx（用它的母片與版面配置）')
    ap.add_argument('--strict', action='store_true', help='檢查有問題就不出檔')
    args = ap.parse_args()

    card, pages = parse(args.outline)
    if not pages:
        sys.exit('%s 裡找不到任何頁 —— 頁標題要寫成「## 3 · 標題」' % args.outline)

    bad, warn = check(card, pages)
    for line in bad:
        sys.stderr.write('⛔ %s\n' % line)
    for line in warn:
        sys.stderr.write('⚠️  %s\n' % line)

    if bad and args.strict:
        sys.exit('\n檢查沒過，沒有出檔。回頭補上面那幾項（--strict 是刻意擋的）。')
    if bad:
        sys.stderr.write('\n上面是閘門該擋的項目。這次仍然出檔了 ——'
                         '要讓它擋下來就加 --strict。\n')

    n = build(card, pages, args.out, args.template)
    sys.stderr.write('\n出了 %d 頁 → %s\n' % (n, args.out))
    sys.stderr.write('用 PowerPoint 開一次，確認文字選得起來、講者備忘有帶進去。\n')
    sys.stderr.write('⚠️ 之後改 pptx 不改大綱 —— 這是單向輸出。\n')


if __name__ == '__main__':
    main()
